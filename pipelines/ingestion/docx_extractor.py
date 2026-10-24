"""Extract paragraphs, headings, and tables from Word documents."""

from io import BytesIO

from docx import Document

from pipelines.ingestion.models import ExtractedPage


def extract_docx(file_bytes: bytes) -> list[ExtractedPage]:
    """Extract DOCX paragraphs and tables with heading metadata."""
    document = Document(BytesIO(file_bytes))
    results: list[ExtractedPage] = []
    heading_stack: list[str] = []
    block_number = 0

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        style_name = paragraph.style.name if paragraph.style else "Normal"
        is_heading = style_name.lower().startswith("heading ")

        if is_heading:
            try:
                level = int(style_name.split()[-1])
            except ValueError:
                level = 1

            level = max(1, level)
            heading_stack = heading_stack[: level - 1]
            heading_stack.append(text)

        results.append(
            ExtractedPage(
                page_number=block_number + 1,
                content=text,
                content_type="text",
                metadata={
                    "source_type": "docx",
                    "block_number": block_number,
                    "style": style_name,
                    "heading_hierarchy": heading_stack.copy(),
                    "is_heading": is_heading,
                },
            )
        )
        block_number += 1

    for table_index, table in enumerate(document.tables):
        rows = [
            [
                cell.text.replace("\n", " ").strip()
                for cell in row.cells
            ]
            for row in table.rows
        ]

        if not rows:
            continue

        markdown = "\n".join(
            "| " + " | ".join(row) + " |"
            for row in rows
        )

        results.append(
            ExtractedPage(
                page_number=block_number + 1,
                content=markdown,
                content_type="table",
                metadata={
                    "source_type": "docx",
                    "table_index": table_index,
                    "heading_hierarchy": heading_stack.copy(),
                },
            )
        )
        block_number += 1

    return results
